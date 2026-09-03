from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
ARTIFACTS.mkdir(parents=True, exist_ok=True)

NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GOLD = "B58A35"
INK = "1F2933"
MUTED = "66727E"
LIGHT = "F2F4F7"
PALE_GOLD = "FBF7ED"
WHITE = "FFFFFF"


def set_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_styles(doc, preset):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10 if preset == "rfi_response" else 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16 if preset == "rfi_response" else 18, 8 if preset == "rfi_response" else 10),
        ("Heading 2", 13, BLUE, 12 if preset == "rfi_response" else 14, 6 if preset == "rfi_response" else 7),
        ("Heading 3", 12, DARK_BLUE, 8 if preset == "rfi_response" else 10, 4 if preset == "rfi_response" else 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5 if preset == "rfi_response" else 0.375)
        style.paragraph_format.first_line_indent = Inches(-0.25 if preset == "rfi_response" else -0.188)
        style.paragraph_format.space_after = Pt(8 if preset == "rfi_response" else 4)
        style.paragraph_format.line_spacing = 1.167 if preset == "rfi_response" else 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run("MAURITANIA AUTOMOTIVE MARKET DEVELOPMENT"), size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    add_page_field(footer)


def add_centerpiece_cover(doc, title, subtitle, metadata):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(56)
    p.paragraph_format.space_after = Pt(12)
    set_font(p.add_run("MARKET DEVELOPMENT DOCUMENT"), size=10, color=GOLD, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run(title), size=25, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    set_font(p.add_run(subtitle), size=13, color=MUTED)

    table = doc.add_table(rows=len(metadata), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2200, 7160])
    for row, (label, value) in zip(table.rows, metadata):
        shade_cell(row.cells[0], LIGHT)
        pr = row.cells[0].paragraphs[0]
        set_font(pr.add_run(label), size=10, color=NAVY, bold=True)
        pv = row.cells[1].paragraphs[0]
        set_font(pv.add_run(value), size=10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("NON-BINDING | SUBJECT TO DUE DILIGENCE AND DEFINITIVE AGREEMENTS"), size=9, color=GOLD, bold=True)
    doc.add_page_break()


def add_partner_cover(doc, title, subtitle, metadata):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run("FIELD EVIDENCE PACK"), size=10, color=GOLD, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run(title), size=28, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    set_font(p.add_run(subtitle), size=13, color=MUTED)

    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [4680, 4680])
    for idx, (label, value) in enumerate(metadata[:2]):
        cell = table.cell(0, idx)
        shade_cell(cell, PALE_GOLD)
        set_font(cell.paragraphs[0].add_run(label.upper()), size=8.5, color=GOLD, bold=True)
        p2 = cell.add_paragraph()
        set_font(p2.add_run(value), size=11, color=NAVY, bold=True)
    for idx, (label, value) in enumerate(metadata[2:4]):
        cell = table.cell(1, idx)
        set_font(cell.paragraphs[0].add_run(label.upper()), size=8.5, color=MUTED, bold=True)
        p2 = cell.add_paragraph()
        set_font(p2.add_run(value), size=10.5, color=INK)
    doc.add_paragraph()


def add_rich_text(paragraph, text, size=11):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        clean = part[2:-2] if bold else part
        set_font(paragraph.add_run(clean), size=size, color=INK, bold=bold)


def add_markdown_body(doc, path, skip_title=True, questionnaire=False):
    lines = path.read_text(encoding="utf-8").splitlines()
    started = not skip_title
    question_count = 0
    pending_category = None
    for raw in lines:
        line = raw.rstrip()
        if skip_title and not started:
            if line.startswith("## "):
                started = True
            else:
                continue
        if not line:
            continue
        if line == "---":
            continue
        if line.startswith("# "):
            continue
        if line.startswith("### "):
            question_text = line[4:]
            is_question = questionnaire and re.match(r"\d+\. ", question_text)
            if is_question:
                next_number = question_count + 1
                if next_number % 2 == 1:
                    doc.add_page_break()
                if pending_category:
                    p = doc.add_paragraph(style="Heading 1")
                    add_rich_text(p, pending_category)
                    pending_category = None
            p = doc.add_paragraph(style="Heading 3")
            add_rich_text(p, question_text)
            if is_question:
                question_count += 1
            continue
        if line.startswith("## "):
            heading_text = line[3:]
            if not questionnaire and heading_text.startswith((
                "3. Requested candidate portfolio",
                "5. Information requested for each model",
                "7. After-sales and parts request",
                "10. Requested response",
                "Contact",
            )):
                doc.add_page_break()
            if questionnaire and re.match(r"^[A-E]\. ", heading_text):
                pending_category = heading_text
                continue
            p = doc.add_paragraph(style="Heading 1")
            add_rich_text(p, heading_text)
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, line[2:])
        elif line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            add_rich_text(p, line)
        else:
            p = doc.add_paragraph()
            add_rich_text(p, line)
            if questionnaire and question_count:
                p.paragraph_format.keep_with_next = True

        if questionnaire and question_count and line and not line.startswith(("#", "-")):
            if re.match(r"^(Which|Who|What|Prepare|Document|Obtain|Identify|Map|Compare)", line):
                add_response_table(doc)


def add_response_table(doc):
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1800, 7560])
    for row, label in zip(table.rows, ("Direct answer", "Evidence / source", "Confirming person", "Gap / next action / date")):
        shade_cell(row.cells[0], LIGHT)
        set_font(row.cells[0].paragraphs[0].add_run(label), size=9, color=NAVY, bold=True)
        p = row.cells[1].paragraphs[0]
        set_font(p.add_run("Response: "), size=9.5, color=MUTED, italic=True)
        p.paragraph_format.space_after = Pt(8)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def build_eoi():
    doc = Document()
    configure_styles(doc, "rfi_response")
    add_centerpiece_cover(
        doc,
        "Non-Binding Expression of Interest",
        "Request for Chinese Automotive Manufacturer and Export Partner Proposals",
        [
            ("Market", "Islamic Republic of Mauritania"),
            ("Prepared by", "Major Dream Williams"),
            ("Submitted through", "Brother Ling"),
            ("Date", "3 September 2026"),
            ("Current phase", "Supplier discovery and local market validation"),
        ],
    )
    add_markdown_body(doc, ROOT / "docs/01-supplier/CHINA-SUPPLIER-EOI-RFP.md")
    out = ARTIFACTS / "Mauritania_China_Automotive_EOI_RFP_2026-09-03.docx"
    doc.save(out)
    return out


def build_sprint():
    doc = Document()
    configure_styles(doc, "compact_reference_guide")
    add_partner_cover(
        doc,
        "Mauritania Automotive Evidence Sprint",
        "Twenty questions that convert market ambition into bankable evidence.",
        [
            ("Prepared for", "Diallo"),
            ("Sprint window", "10-14 days from acceptance"),
            ("Prepared by", "Major Dream Williams / Program Coordination"),
            ("Status", "Field diligence - no commercial rights granted"),
        ],
    )
    add_markdown_body(doc, ROOT / "docs/02-field/DIALLO-20-QUESTION-EVIDENCE-SPRINT.md", questionnaire=True)
    out = ARTIFACTS / "Diallo_Mauritania_Automotive_20_Question_Evidence_Sprint.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    for result in (build_eoi(), build_sprint()):
        print(result)
